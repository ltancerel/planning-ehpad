#!/usr/bin/env python3
"""
Outil d'import en masse pour planning-ehpad-prod (Supabase) — tourne sur votre
poste/serveur, aucune donnée n'est envoyée ailleurs que vers votre propre base.

Reprend et étend le script d'import de salariés du 28/10 : remplace la notion
de contrat (CDD/CDI, table contrat) par le champ type_paie (25/M, 30/M, Aucun
— migration du 29/10), et ajoute un deuxième mode d'import pour les codes
horaires (même format que le modèle Excel livré le 28/10).

Par défaut (ou avec --test), le script ne fait qu'un test à blanc : il lit le
ou les fichiers, valide toutes les lignes et affiche un rapport, sans rien
écrire en base. Rien n'est écrit tant que --import n'est pas explicitement
passé. Les erreurs de salarié n'affichent jamais de nom ni de prénom, pour
que le rapport puisse être partagé sans risque — seulement des numéros de
ligne Excel.

Usage :
  export SUPABASE_SERVICE_ROLE_KEY='...'   # clé "service_role" (Settings > API)

  # test à blanc (par défaut)
  python3 planning-ehpad-tool.py --ehpad "Nom de l'EHPAD" --salaries salaries.xlsx
  python3 planning-ehpad-tool.py --ehpad "Nom de l'EHPAD" --codes-horaires codes.xlsx
  python3 planning-ehpad-tool.py --ehpad "Nom de l'EHPAD" --salaries s.xlsx --codes-horaires c.xlsx

  # import réel (les deux options ci-dessus peuvent aussi prendre --import)
  python3 planning-ehpad-tool.py --ehpad "Nom de l'EHPAD" --salaries salaries.xlsx --import

  # inspection des en-têtes d'un fichier (aucune écriture, pas de clé requise)
  python3 planning-ehpad-tool.py --salaries salaries.xlsx --entetes

Dépendances : pip3 install openpyxl requests
"""

import argparse
import os
import re
import sys
import unicodedata
from datetime import date, datetime

import requests
from openpyxl import load_workbook

SUPABASE_URL = "https://xacgsqvuxjungpxpmqmz.supabase.co"


# =====================================================================
# Utilitaires communs
# =====================================================================

def normaliser(valeur):
    """Minuscules, sans accents ni espaces superflus, pour comparer des intitulés/valeurs."""
    valeur = unicodedata.normalize("NFKD", str(valeur or "")).encode("ascii", "ignore").decode()
    return " ".join(valeur.lower().split())


def texte(valeur):
    if valeur is None:
        return None
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    valeur = str(valeur).strip()
    return valeur or None


def en_date(valeur):
    if valeur in (None, ""):
        return None
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, date):
        return valeur.isoformat()
    valeur = str(valeur).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(valeur, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"date illisible « {valeur} » (attendu JJ/MM/AAAA)")


def en_heure(valeur):
    """Accepte un datetime.time (Excel) ou une chaîne HH:MM / HH:MM:SS."""
    if valeur is None or valeur == "":
        return None
    if hasattr(valeur, "strftime"):
        return valeur.strftime("%H:%M")
    valeur = str(valeur).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(valeur, fmt).strftime("%H:%M")
        except ValueError:
            pass
    raise ValueError(f"heure illisible « {valeur} » (attendu HH:MM)")


def en_booleen(valeur, defaut=False):
    v = normaliser(valeur)
    if v in ("", ) and defaut is not None:
        return defaut
    if v in ("vrai", "true", "oui", "1", "x"):
        return True
    if v in ("faux", "false", "non", "0"):
        return False
    raise ValueError(f"valeur booléenne « {valeur} » invalide (VRAI ou FAUX)")


class Supabase:
    def __init__(self, cle):
        self.session = requests.Session()
        self.session.headers.update({
            "apikey": cle,
            "Authorization": f"Bearer {cle}",
            "Content-Type": "application/json",
        })

    def _verifier(self, reponse):
        if reponse.ok:
            return reponse.json() if reponse.content else None
        # On n'affiche que le code et le message : le champ "details" de Postgres
        # peut recopier la ligne refusée, donc potentiellement des noms.
        try:
            erreur = reponse.json()
            message = f"{erreur.get('code', reponse.status_code)} {erreur.get('message', '')}"
        except ValueError:
            message = f"HTTP {reponse.status_code}"
        raise RuntimeError(message)

    def lire(self, table, **filtres):
        return self._verifier(self.session.get(f"{SUPABASE_URL}/rest/v1/{table}", params=filtres))

    def inserer(self, table, ligne):
        r = self.session.post(
            f"{SUPABASE_URL}/rest/v1/{table}",
            json=ligne,
            params={"select": "id"},
            headers={"Prefer": "return=representation"},
        )
        return self._verifier(r)[0]["id"]

    def supprimer(self, table, **filtres):
        params = {cle: f"eq.{valeur}" for cle, valeur in filtres.items()}
        self._verifier(self.session.delete(f"{SUPABASE_URL}/rest/v1/{table}", params=params))


def lire_excel(chemin, feuille=None):
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    ws = classeur[feuille] if feuille else classeur.active
    lignes = ws.iter_rows(values_only=True)
    entetes = [texte(c) or "" for c in next(lignes)]
    return entetes, [(i, ligne) for i, ligne in enumerate(lignes, start=2)]


def indexer_colonnes(entetes, colonnes, obligatoires, nom_fichier):
    index_entetes = {normaliser(e): i for i, e in enumerate(entetes) if e}
    col = {}
    for champ, intitule in colonnes.items():
        if normaliser(intitule) in index_entetes:
            col[champ] = index_entetes[normaliser(intitule)]
        elif champ in obligatoires:
            sys.exit(
                f"[{nom_fichier}] Colonne obligatoire « {intitule} » introuvable. "
                f"Lancez avec --entetes pour voir les en-têtes détectés."
            )
    return col


# =====================================================================
# Import des salariés
# =====================================================================

# Correspondance champ de l'application -> intitulé de colonne attendu dans
# l'Excel. "Contrat Actif"/"Type de contrat" (ancien modèle contrat,
# supprimé le 29/10) et "Manager" (étiquette figée sans lien avec un vrai
# compte, retirée le 30/10 — le vrai responsable se lit via
# service.responsable_id, écran Utilisateurs) ne sont plus reconnues : une
# colonne de ce nom encore présente dans un fichier client est
# silencieusement ignorée.
COLONNES_SALARIES = {
    "matricule": "Matricule",
    "nom": "Nom",
    "prenom": "Prénom",
    "service": "Service",
    "type_paie": "Type de paie",
    "alignement_roulement": "Alignement roulement",
    "presence": "Présence",
    "roulement": "Roulement",
    "roulement_debut": "Date début roulement",
}
OBLIGATOIRES_SALARIES = ["matricule", "nom", "prenom", "service"]

TYPES_PAIE_VALIDES = {"Aucun", "25/M", "30/M"}
GROUPES_ROULEMENT_VALIDES = {"Équipe A", "Équipe B", "Équipe C", "Équipe D"}
# Certains fichiers client expriment l'alignement roulement en Paire/Impaire
# plutôt qu'en 4 équipes — correspondance actée avec le client le 29/10.
ALIGNEMENT_PAIRE_IMPAIRE = {"Paire": "Équipe A", "Impaire": "Équipe B"}


def en_presence(valeur):
    v = normaliser(valeur)
    if v in ("", "present", "p", "oui"):
        return "Présent"
    if v in ("absent", "a", "non"):
        return "Absent"
    raise ValueError(f"présence « {valeur} » invalide (Présent ou Absent)")


def en_type_paie(valeur):
    v = texte(valeur)
    if v is None or normaliser(v) in ("", "aucun"):
        return "Aucun"
    for t in TYPES_PAIE_VALIDES:
        if normaliser(t) == normaliser(v):
            return t
    raise ValueError(f"type de paie « {valeur} » invalide ({', '.join(sorted(TYPES_PAIE_VALIDES))})")


def en_alignement_roulement(valeur):
    v = texte(valeur)
    if v is None:
        return "Équipe A"  # défaut applicatif, cf. SalarieForm.tsx
    if v in ALIGNEMENT_PAIRE_IMPAIRE:
        return ALIGNEMENT_PAIRE_IMPAIRE[v]
    for g in GROUPES_ROULEMENT_VALIDES:
        if normaliser(g) == normaliser(v):
            return g
    raise ValueError(
        f"alignement roulement « {valeur} » invalide "
        f"(Équipe A à D, ou Paire/Impaire)"
    )


def en_matricule(valeur):
    v = (texte(valeur) or "").upper()
    if len(v) != 4 or not v.isalpha():
        raise ValueError(f"matricule « {valeur} » invalide (4 lettres exactement)")
    return v


def preparer_import_salaries(db, chemin, feuille, ehpad_id):
    entetes, lignes = lire_excel(chemin, feuille)
    col = indexer_colonnes(entetes, COLONNES_SALARIES, OBLIGATOIRES_SALARIES, os.path.basename(chemin))

    services = {normaliser(s["nom"]): s["id"]
                for s in db.lire("service", select="id,nom", ehpad_id=f"eq.{ehpad_id}")}
    roulements = {normaliser(r["nom"]): r["id"]
                  for r in db.lire("roulement", select="id,nom", ehpad_id=f"eq.{ehpad_id}")}
    matricules_existants = {normaliser(s["matricule"])
                             for s in db.lire("salarie", select="matricule", ehpad_id=f"eq.{ehpad_id}")}

    def val(ligne, champ):
        i = col.get(champ)
        return ligne[i] if i is not None and i < len(ligne) else None

    a_importer, erreurs, deja_presents, vides = [], [], 0, 0
    vus = set()
    for num, ligne in lignes:
        if all(c in (None, "") for c in ligne):
            vides += 1
            continue
        try:
            matricule = en_matricule(val(ligne, "matricule"))
            nom, prenom = texte(val(ligne, "nom")), texte(val(ligne, "prenom"))
            if not (nom and prenom):
                raise ValueError("nom ou prénom vide")

            if normaliser(matricule) in matricules_existants:
                deja_presents += 1
                continue
            if normaliser(matricule) in vus:
                raise ValueError(f"matricule {matricule} en double dans le fichier")
            vus.add(normaliser(matricule))

            service = texte(val(ligne, "service"))
            if normaliser(service) not in services:
                raise ValueError(f"service « {service} » inconnu dans l'application")

            salarie = {
                "ehpad_id": ehpad_id,
                "service_id": services[normaliser(service)],
                "matricule": matricule,
                "nom": nom.upper(),
                "prenom": prenom[:1].upper() + prenom[1:].lower(),
                "alignement_roulement": en_alignement_roulement(val(ligne, "alignement_roulement")),
                "presence": en_presence(val(ligne, "presence")),
                "type_paie": en_type_paie(val(ligne, "type_paie")),
            }

            affectation = None
            roulement = texte(val(ligne, "roulement"))
            if roulement:
                if normaliser(roulement) not in roulements:
                    raise ValueError(f"roulement « {roulement} » inconnu dans l'application")
                debut = en_date(val(ligne, "roulement_debut"))
                if not debut:
                    raise ValueError("date de début de roulement manquante")
                affectation = {"roulement_id": roulements[normaliser(roulement)], "date_debut": debut}

            a_importer.append((num, salarie, affectation))
        except ValueError as e:
            erreurs.append((num, str(e)))

    rapport = [
        f"[{os.path.basename(chemin)}] salariés — à importer : {len(a_importer)}, "
        f"déjà présents : {deja_presents}, lignes vides : {vides}, erreurs : {len(erreurs)}",
    ]
    for num, msg in erreurs:
        rapport.append(f"    ligne {num} : {msg}")
    return a_importer, erreurs, rapport


def executer_import_salaries(db, ehpad_id, a_importer):
    importes, echecs = 0, []
    for num, salarie, affectation in a_importer:
        salarie_id = None
        try:
            salarie_id = db.inserer("salarie", salarie)
            if affectation:
                db.inserer("affectation_roulement", {**affectation, "salarie_id": salarie_id})
            importes += 1
        except RuntimeError as e:
            echecs.append((num, str(e)))
            if salarie_id:
                db.supprimer("affectation_roulement", salarie_id=salarie_id)
                db.supprimer("salarie", id=salarie_id)

    _journaliser(db, ehpad_id, "salarie", importes, echecs,
                 f"Import Excel salariés : {importes} créé(s), {len(echecs)} échec(s)")
    return importes, echecs


# =====================================================================
# Import des codes horaires (même format que le modèle Excel du 28/10)
# =====================================================================

COLONNES_CODES_HORAIRES = {
    "code": "code",
    "intitule": "intitule",
    "categorie": "categorie",
    "couleur_fond": "couleur_fond",
    "couleur_texte": "couleur_texte",
    "afficher_vue_annuelle": "afficher_vue_annuelle",
    "commentaire": "commentaire",
    "plage1_debut": "plage1_debut",
    "plage1_fin": "plage1_fin",
    "plage2_debut": "plage2_debut",
    "plage2_fin": "plage2_fin",
    "plage3_debut": "plage3_debut",
    "plage3_fin": "plage3_fin",
    "plage4_debut": "plage4_debut",
    "plage4_fin": "plage4_fin",
    "action": "action",
    "type_evenement": "type_evenement",
    "duree_heures": "duree_heures",
}
OBLIGATOIRES_CODES_HORAIRES = ["code", "intitule", "categorie", "couleur_fond", "couleur_texte"]

CATEGORIES_VALIDES = {"travail", "informatif", "evenementiel"}
TYPES_EVENEMENT_VALIDES = {"special", "normal", "partiel"}
REGEX_COULEUR = re.compile(r"^#[0-9A-Fa-f]{6}$")


def en_couleur(valeur, nom_champ):
    v = texte(valeur)
    if not v or not REGEX_COULEUR.match(v):
        raise ValueError(f"{nom_champ} « {valeur} » invalide (format #RRGGBB attendu)")
    return v


def en_categorie(valeur):
    v = normaliser(valeur)
    for c in CATEGORIES_VALIDES:
        if normaliser(c) == v:
            return c
    raise ValueError(f"categorie « {valeur} » invalide ({', '.join(sorted(CATEGORIES_VALIDES))})")


def en_type_evenement(valeur):
    v = normaliser(valeur)
    for t in TYPES_EVENEMENT_VALIDES:
        if normaliser(t) == v:
            return t
    raise ValueError(f"type_evenement « {valeur} » invalide ({', '.join(sorted(TYPES_EVENEMENT_VALIDES))})")


def preparer_import_codes_horaires(db, chemin, feuille, ehpad_id):
    entetes, lignes = lire_excel(chemin, feuille)
    col = indexer_colonnes(entetes, COLONNES_CODES_HORAIRES, OBLIGATOIRES_CODES_HORAIRES, os.path.basename(chemin))

    codes_existants = {normaliser(c["code"])
                        for c in db.lire("code_horaire", select="code", ehpad_id=f"eq.{ehpad_id}")}

    def val(ligne, champ):
        i = col.get(champ)
        return ligne[i] if i is not None and i < len(ligne) else None

    a_importer, erreurs, deja_presents, vides = [], [], 0, 0
    vus = set()
    for num, ligne in lignes:
        if all(c in (None, "") for c in ligne):
            vides += 1
            continue
        try:
            code = texte(val(ligne, "code"))
            intitule = texte(val(ligne, "intitule"))
            if not code or not intitule:
                raise ValueError("code ou intitulé vide")

            if normaliser(code) in codes_existants:
                deja_presents += 1
                continue
            if normaliser(code) in vus:
                raise ValueError(f"code {code} en double dans le fichier")
            vus.add(normaliser(code))

            categorie = en_categorie(val(ligne, "categorie"))
            code_horaire = {
                "ehpad_id": ehpad_id,
                "code": code,
                "intitule": intitule,
                "categorie": categorie,
                "couleur_fond": en_couleur(val(ligne, "couleur_fond"), "couleur_fond"),
                "couleur_texte": en_couleur(val(ligne, "couleur_texte"), "couleur_texte"),
                "afficher_vue_annuelle": en_booleen(val(ligne, "afficher_vue_annuelle"), defaut=False),
                "commentaire": texte(val(ligne, "commentaire")),
            }

            # Plages horaires — categorie = travail uniquement, jusqu'à 4,
            # remplies consécutivement à partir de plage1 (même règle que le
            # formulaire Admin > Codes horaires).
            plages = []
            for i in (1, 2, 3, 4):
                debut = en_heure(val(ligne, f"plage{i}_debut"))
                fin = en_heure(val(ligne, f"plage{i}_fin"))
                if debut is None and fin is None:
                    break
                if debut is None or fin is None:
                    raise ValueError(f"plage{i} incomplète (début et fin tous les deux requis)")
                plages.append({"ordre": i, "heure_debut": debut, "heure_fin": fin})
            if plages and categorie != "travail":
                raise ValueError("des plages horaires sont renseignées pour une categorie autre que travail")

            # Événementiel/action/type_evenement/duree_heures — mêmes règles
            # que les CHECK posés en base (migration domaine C).
            action = texte(val(ligne, "action"))
            type_evenement_brut = texte(val(ligne, "type_evenement"))
            duree_heures = texte(val(ligne, "duree_heures"))

            if categorie == "evenementiel":
                if not type_evenement_brut:
                    raise ValueError("type_evenement obligatoire pour categorie = evenementiel")
                type_evenement = en_type_evenement(type_evenement_brut)
                if type_evenement == "special" and duree_heures:
                    raise ValueError("duree_heures doit rester vide pour type_evenement = special")
                if type_evenement == "normal" and not duree_heures:
                    raise ValueError("duree_heures obligatoire pour type_evenement = normal")
            else:
                if type_evenement_brut or duree_heures or action:
                    raise ValueError(
                        "action/type_evenement/duree_heures doivent rester vides "
                        "hors categorie = evenementiel"
                    )
                type_evenement = None

            code_horaire["action"] = action
            code_horaire["type_evenement"] = type_evenement
            code_horaire["duree_heures"] = float(duree_heures) if duree_heures else None

            a_importer.append((num, code_horaire, plages))
        except ValueError as e:
            erreurs.append((num, str(e)))

    rapport = [
        f"[{os.path.basename(chemin)}] codes horaires — à importer : {len(a_importer)}, "
        f"déjà présents : {deja_presents}, lignes vides : {vides}, erreurs : {len(erreurs)}",
    ]
    for num, msg in erreurs:
        rapport.append(f"    ligne {num} : {msg}")
    return a_importer, erreurs, rapport


def executer_import_codes_horaires(db, ehpad_id, a_importer):
    importes, echecs = 0, []
    for num, code_horaire, plages in a_importer:
        code_horaire_id = None
        try:
            code_horaire_id = db.inserer("code_horaire", code_horaire)
            for plage in plages:
                db.inserer("plage_horaire", {**plage, "code_horaire_id": code_horaire_id})
            importes += 1
        except RuntimeError as e:
            echecs.append((num, str(e)))
            if code_horaire_id:
                db.supprimer("plage_horaire", code_horaire_id=code_horaire_id)
                db.supprimer("code_horaire", id=code_horaire_id)

    _journaliser(db, ehpad_id, "code_horaire", importes, echecs,
                 f"Import Excel codes horaires : {importes} créé(s), {len(echecs)} échec(s)")
    return importes, echecs


def _journaliser(db, ehpad_id, entite, importes, echecs, message):
    try:
        db.inserer("log_audit", {
            "ehpad_id": ehpad_id,
            "acteur_type": "service",
            "acteur_nom": "planning-ehpad-tool.py",
            "niveau": "avertissement" if echecs else "info",
            "statut": "echec" if echecs and not importes else "succes",
            "type_operation": "import",
            "entite": entite,
            "source": "script",
            "message": message,
        })
    except RuntimeError as e:
        print(f"(journal d'audit non écrit : {e})")


# =====================================================================
# Point d'entrée
# =====================================================================

def main():
    p = argparse.ArgumentParser(
        description="Import en masse (salariés et/ou codes horaires) dans planning-ehpad-prod",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    p.add_argument("--ehpad", help="nom de l'EHPAD tel qu'enregistré dans l'application")
    p.add_argument("--salaries", metavar="FICHIER", help="fichier Excel des salariés à importer")
    p.add_argument("--codes-horaires", metavar="FICHIER", help="fichier Excel des codes horaires à importer")
    p.add_argument("--feuille", help="nom de l'onglet à lire (par défaut : l'onglet actif de chaque fichier)")
    p.add_argument("--entetes", action="store_true",
                    help="affiche seulement les en-têtes des fichiers fournis et quitte (pas d'accès base)")
    mode = p.add_mutually_exclusive_group()
    mode.add_argument("--test", action="store_true", help="test à blanc, rien n'est écrit (comportement par défaut)")
    mode.add_argument("--import", dest="executer", action="store_true", help="écrit réellement dans la base")
    args = p.parse_args()

    if not args.salaries and not args.codes_horaires:
        sys.exit("Donnez au moins --salaries ou --codes-horaires.")

    if args.entetes:
        for chemin in (args.salaries, args.codes_horaires):
            if not chemin:
                continue
            entetes, _ = lire_excel(chemin, args.feuille)
            print(f"[{os.path.basename(chemin)}] en-têtes trouvés :")
            for e in entetes:
                print(f"  - {e}")
        return

    if not args.ehpad:
        sys.exit("--ehpad est requis (sauf avec --entetes).")

    cle = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not cle:
        sys.exit("Variable SUPABASE_SERVICE_ROLE_KEY non définie.")
    db = Supabase(cle)

    ehpads = [e for e in db.lire("ehpad", select="id,nom") if normaliser(e["nom"]) == normaliser(args.ehpad)]
    if len(ehpads) != 1:
        sys.exit(f"EHPAD « {args.ehpad} » introuvable (ou en double) dans la base.")
    ehpad_id = ehpads[0]["id"]

    rapport_complet = []
    a_importer_salaries = erreurs_salaries = None
    a_importer_codes = erreurs_codes = None

    if args.salaries:
        a_importer_salaries, erreurs_salaries, rapport = preparer_import_salaries(
            db, args.salaries, args.feuille, ehpad_id)
        rapport_complet += rapport
    if args.codes_horaires:
        a_importer_codes, erreurs_codes, rapport = preparer_import_codes_horaires(
            db, args.codes_horaires, args.feuille, ehpad_id)
        rapport_complet += rapport

    print("\n".join(rapport_complet))

    y_a_t_il_des_erreurs = bool(erreurs_salaries) or bool(erreurs_codes)
    if y_a_t_il_des_erreurs:
        sys.exit("\nCorrigez les erreurs ci-dessus puis relancez. Rien n'a été écrit.")

    if not args.executer:
        print("\nTest à blanc terminé, rien n'a été écrit. Relancez avec --import pour importer réellement.")
        return

    if a_importer_salaries is not None:
        importes, echecs = executer_import_salaries(db, ehpad_id, a_importer_salaries)
        print(f"\nSalariés importés : {importes}, échecs : {len(echecs)}")
        for num, msg in echecs:
            print(f"  ligne {num} : {msg}")
    if a_importer_codes is not None:
        importes, echecs = executer_import_codes_horaires(db, ehpad_id, a_importer_codes)
        print(f"\nCodes horaires importés : {importes}, échecs : {len(echecs)}")
        for num, msg in echecs:
            print(f"  ligne {num} : {msg}")


if __name__ == "__main__":
    main()
